#!/usr/bin/env python3
"""Normalize NMSB collection, dataset, and resolved eLab metadata for LinkML."""

from __future__ import annotations

import argparse
import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, Iterable, Optional, Tuple

import yaml
from openpyxl import load_workbook


NA_VALUES = {"", "na", "n/a", "none", "null", "not applicable"}

COLLECTION_ENUMS = {
    "nmsb collection": "NMSBCollection",
    "nmsbcollection": "NMSBCollection",
}
DATA_CATEGORY_ENUMS = {"characterization": "characterization"}
CORE_CAPABILITY_ENUMS = {"microbial measurements": "microbial_measurements"}
PRIMARY_FOCUS_ENUMS = {"microbiome": "microbiome"}
DIVISION_ENUMS = {
    "biosystems and biomaterials division, division 644": (
        "644_biosystems_and_biomaterials_division"
    )
}


def _clean_scalar(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return str(value)
    text = str(value).strip()
    return None if text.lower() in NA_VALUES else text


def read_key_value_workbook(path: Path) -> Dict[str, Any]:
    """Read the first worksheet's field/value columns into a dictionary."""
    workbook = load_workbook(path, data_only=True, read_only=True)
    worksheet = workbook.worksheets[0]
    values: Dict[str, Any] = {}
    for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        if row_number == 1 or not row:
            continue
        key = _clean_scalar(row[0])
        if not key:
            continue
        values[str(key)] = _clean_scalar(row[1] if len(row) > 1 else None)
    workbook.close()
    return values


def _enum(value: Any, choices: Dict[str, str]) -> Any:
    cleaned = _clean_scalar(value)
    if cleaned is None:
        return None
    return choices.get(str(cleaned).lower(), cleaned)


def _replace_nist_id(value: Any, nist_id: str) -> Any:
    cleaned = _clean_scalar(value)
    if cleaned is None:
        return None
    return re.sub(r"\bNIST\d+\b", nist_id, str(cleaned), flags=re.IGNORECASE)


def _put(payload: Dict[str, Any], key: str, value: Any) -> None:
    cleaned = _clean_scalar(value)
    if cleaned is not None:
        payload[key] = cleaned


def normalize_nmsb(
    collection: Dict[str, Any],
    dataset: Dict[str, Any],
    elab: Dict[str, Any],
    nist_id: str,
    dataset_id: str,
) -> Tuple[Dict[str, Any], Dict[str, Any]]:
    """Build the validation payload and an audit report for the applied decisions."""
    source_title = _clean_scalar(dataset.get("dataset_title"))
    payload: Dict[str, Any] = {
        "collection_name": _enum(
            collection.get("collection_name"), COLLECTION_ENUMS
        ),
        "principal_contact_name": _clean_scalar(
            collection.get("principal_contact_name")
        ),
        "data_category": _enum(
            collection.get("data_category"), DATA_CATEGORY_ENUMS
        ),
        "participating_NIST_division": [
            _enum(collection.get("participating_NIST_division"), DIVISION_ENUMS)
        ],
        "dataset_title": _replace_nist_id(dataset.get("dataset_title"), nist_id),
        "dataset_description": _replace_nist_id(
            dataset.get("dataset_description"), nist_id
        ),
        "dataset_ID": dataset_id,
    }

    for key in (
        "collection_description",
        "principal_contact_email",
        "data_custodian_name",
        "data_custodian_email",
        "associated_consortium",
        "owner_principal",
        "URL_link",
        "URL_link_description",
    ):
        _put(payload, key, collection.get(key))

    core_capability = _enum(
        collection.get("core_capabilities"), CORE_CAPABILITY_ENUMS
    )
    if core_capability:
        payload["core_capabilities"] = [core_capability]
    primary_focus = _enum(
        collection.get("primary_focus_areas"), PRIMARY_FOCUS_ENUMS
    )
    if primary_focus:
        payload["primary_focus_areas"] = [primary_focus]

    dataset_values = {
        "external_organization_name": dataset.get("external_organization_name"),
        "external_organization_ID": dataset.get("external_organization_ID"),
        "external_organization_URI": dataset.get("external_organization_URI"),
        "data_custodian_name": dataset.get("data_custodian_name"),
        "data_custodian_email": dataset.get("data_custodian_email"),
        "assay_technique": dataset.get("assay_technique"),
        "material_type": dataset.get("material_type"),
        "starting_material_type": dataset.get("starting_material_type"),
        "data_capture_start_date": dataset.get("data_capture_start_date"),
        "data_capture_end_date": dataset.get("data_capture_end_date"),
        "original_upload_location": dataset.get("original_upload_location"),
        "derived_from_dataset_ID": dataset.get("derived_from_dataset_ID"),
        "data_submitter": dataset.get("data_submitter"),
        "dataset_creation_date": dataset.get("dataset_creation_date"),
        "dataset_modified_date": dataset.get("dataset_modified_date"),
        "workflow_id": dataset.get("workflow_id"),
        "workflow_type": dataset.get("workflow_type"),
        "software_tool_name": dataset.get("software_tool_name"),
        "software_version": dataset.get("software_version"),
        "execution_date": dataset.get("execution_date"),
        "execution_environment": dataset.get("execution_environment"),
        "executed_by": dataset.get("executed_by"),
        "command_executed": dataset.get("command_executed"),
        "parameters_used": dataset.get("parameters_used"),
        "workflow_notes": dataset.get("workflow_notes"),
    }
    for key, value in dataset_values.items():
        _put(payload, key, value)

    resolved_fields = elab.get("resolved_fields", {})
    if not isinstance(resolved_fields, dict):
        raise ValueError("eLab sidecar must contain an object named 'resolved_fields'")
    for key, value in resolved_fields.items():
        if isinstance(value, list):
            cleaned_values = []
            for item in value:
                cleaned = _clean_scalar(item)
                if cleaned is not None:
                    cleaned_values.append(cleaned)
            if cleaned_values:
                payload[key] = cleaned_values
        else:
            _put(payload, key, value)

    payload["material_reference_code"] = nist_id
    payload["NIST_ID"] = nist_id
    supplier = _clean_scalar(payload.get("material_supplier_institution"))
    lot_code = _clean_scalar(payload.get("material_lot_code"))
    title_parts = [part for part in (supplier, nist_id, lot_code) if part]
    payload["elabs_resource_title"] = "_".join(str(part) for part in title_parts)

    report = {
        "source_collection_workbook_value": collection.get("collection_name"),
        "source_dataset_title": source_title,
        "normalized_nist_id": nist_id,
        "generated_dataset_ID": dataset_id,
        "elabs_resource_title": payload["elabs_resource_title"],
        "excluded_unresolved_fields": elab.get("excluded_unresolved_fields", {}),
        "source_files": elab.get("source_files", []),
    }
    return payload, report


def write_outputs(
    payload: Dict[str, Any], report: Dict[str, Any], output: Path, report_path: Path
) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(yaml.safe_dump(payload, sort_keys=False), encoding="utf-8")
    report_path.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")


def parse_args(argv: Optional[Iterable[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Normalize an NMSB package for NMSBCollection validation."
    )
    parser.add_argument("--collection-workbook", required=True, type=Path)
    parser.add_argument("--dataset-workbook", required=True, type=Path)
    parser.add_argument("--elab-json", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--report", required=True, type=Path)
    parser.add_argument("--nist-id", required=True)
    parser.add_argument("--dataset-id", required=True)
    return parser.parse_args(argv)


def main() -> int:
    args = parse_args()
    collection = read_key_value_workbook(args.collection_workbook)
    dataset = read_key_value_workbook(args.dataset_workbook)
    elab = json.loads(args.elab_json.read_text(encoding="utf-8"))
    payload, report = normalize_nmsb(
        collection=collection,
        dataset=dataset,
        elab=elab,
        nist_id=args.nist_id,
        dataset_id=args.dataset_id,
    )
    write_outputs(payload, report, args.output, args.report)
    print(f"Wrote normalized NMSB payload to {args.output}")
    print(f"Wrote normalization report to {args.report}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
